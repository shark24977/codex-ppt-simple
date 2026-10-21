from __future__ import annotations

import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
EXAMPLE = ROOT / "assets" / "deck_spec.example.json"
SEMANTIC_EXAMPLE = ROOT / "assets" / "deck_spec.semantic.example.json"


class EditableDeckSmokeTest(unittest.TestCase):
    def test_example_builds_with_native_objects(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "editable-example.pptx"
            result = subprocess.run(
                [sys.executable, str(SCRIPTS / "build_editable_ppt.py"), str(EXAMPLE), "--out", str(out)],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertTrue(out.is_file())
            prs = Presentation(str(out))
            self.assertEqual(len(prs.slides), 5)
            self.assertTrue(any(getattr(shape, "has_chart", False) for shape in prs.slides[3].shapes))
            for slide in prs.slides:
                self.assertTrue(
                    any(
                        getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip()
                        for shape in slide.shapes
                    )
                )

            qa = subprocess.run(
                [sys.executable, str(SCRIPTS / "qa_editability.py"), str(out), "--strict"],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertEqual(qa.returncode, 0, qa.stdout + qa.stderr)
            report = json.loads(qa.stdout)
            self.assertTrue(report["editable"])
            self.assertEqual(report["totals"]["full_slide_pictures"], 0)
            self.assertGreater(report["totals"]["native_text_shapes"], 0)

    def test_readability_qa_passes_for_generated_example(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "readable-example.pptx"
            build = subprocess.run(
                [sys.executable, str(SCRIPTS / "build_editable_ppt.py"), str(EXAMPLE), "--out", str(out)],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertEqual(build.returncode, 0, build.stdout + build.stderr)
            qa = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPTS / "qa_readability.py"),
                    str(out),
                    "--spec",
                    str(EXAMPLE),
                    "--strict",
                ],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertEqual(qa.returncode, 0, qa.stdout + qa.stderr)
            report = json.loads(qa.stdout)
            self.assertTrue(report["passed"])
            self.assertEqual(report["totals"]["undersized_runs"], 0)

    def test_inline_emphasis_becomes_native_bold_runs(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            spec_path = tmp_path / "emphasis.json"
            out = tmp_path / "emphasis.pptx"
            spec_path.write_text(
                json.dumps(
                    {
                        "version": 1,
                        "mode": "native",
                        "slides": [
                            {
                                "type": "title_body",
                                "title": "重点应形成原生文字层级",
                                "takeaway": "结论优先于说明。",
                                "bullets": ["**关键机制：**联合分析语义与硬件约束。"],
                            }
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            build = subprocess.run(
                [sys.executable, str(SCRIPTS / "build_editable_ppt.py"), str(spec_path), "--out", str(out)],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertEqual(build.returncode, 0, build.stdout + build.stderr)
            prs = Presentation(str(out))
            bullet_shape = next(shape for shape in prs.slides[0].shapes if shape.name == "Editable bullets")
            self.assertNotIn("**", bullet_shape.text)
            runs = bullet_shape.text_frame.paragraphs[0].runs
            self.assertTrue(any(run.font.bold and "关键机制" in run.text for run in runs))
            self.assertTrue(all(run.font.size.pt >= 13.5 for run in runs if run.text.strip()))

    def test_readability_qa_rejects_small_unemphasized_body(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "too-small.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            title = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(11.5), Inches(0.6))
            title.name = "Editable slide title"
            title.text_frame.text = "小字号不应通过成品门禁"
            title_run = title.text_frame.paragraphs[0].runs[0]
            title_run.font.size = Pt(28)
            title_run.font.bold = True
            body = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11), Inches(4))
            body.name = "Editable body"
            body.text_frame.text = "这是一段没有重点层级的正文。" * 6
            body_run = body.text_frame.paragraphs[0].runs[0]
            body_run.font.size = Pt(9)
            prs.save(out)
            qa = subprocess.run(
                [sys.executable, str(SCRIPTS / "qa_readability.py"), str(out), "--strict"],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertNotEqual(qa.returncode, 0)
            report = json.loads(qa.stdout)
            self.assertFalse(report["passed"])
            self.assertGreater(report["totals"]["undersized_runs"], 0)
            self.assertTrue(any("bold emphasis ratio" in warning for warning in report["warnings"]))

    def test_reference_theme_path_extracts_and_rebuilds(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            reference = tmp_path / "reference.pptx"
            profile = tmp_path / "style.json"
            output = tmp_path / "from-reference.pptx"
            spec_path = tmp_path / "deck_spec.json"

            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            accent = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(0.7),
                Inches(0.7),
                Inches(2.0),
                Inches(0.25),
            )
            accent.fill.solid()
            accent.fill.fore_color.rgb = RGBColor(0x12, 0x78, 0x7A)
            accent.line.fill.background()
            textbox = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(6.0), Inches(1.0))
            textbox.text_frame.text = "Reference style"
            run = textbox.text_frame.paragraphs[0].runs[0]
            run.font.name = "Aptos Display"
            run.font.size = Pt(28)
            run.font.color.rgb = RGBColor(0x17, 0x21, 0x2B)
            prs.save(reference)

            extract = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPTS / "extract_reference_style.py"),
                    str(reference),
                    "--out",
                    str(profile),
                ],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertEqual(extract.returncode, 0, extract.stdout + extract.stderr)
            extracted = json.loads(profile.read_text(encoding="utf-8"))
            self.assertEqual(extracted["slides_analyzed"], 1)
            self.assertIn("colors", extracted)

            spec_path.write_text(
                json.dumps(
                    {
                        "version": 1,
                        "title": "Reference rebuild",
                        "mode": "native",
                        "slides": [
                            {"type": "cover", "title": "Editable cover", "subtitle": "Preserved theme"},
                            {"type": "closing", "title": "Editable closing", "subtitle": "Native objects"},
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            build = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPTS / "build_editable_ppt.py"),
                    str(spec_path),
                    "--reference-pptx",
                    str(reference),
                    "--style-profile",
                    str(profile),
                    "--out",
                    str(output),
                ],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertEqual(build.returncode, 0, build.stdout + build.stderr)
            rebuilt = Presentation(str(output))
            self.assertEqual(len(rebuilt.slides), 2)
            self.assertEqual(len(rebuilt.slide_masters), len(prs.slide_masters))

    def test_reference_slide_family_prefers_local_palette(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            reference = tmp_path / "families.pptx"
            profile = tmp_path / "local-style.json"
            prs = Presentation()
            for color in (
                RGBColor(0x44, 0x72, 0xC4),
                RGBColor(0xB2, 0x1F, 0x2D),
                RGBColor(0x17, 0x3F, 0x5F),
            ):
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                for offset in (0.0, 0.5):
                    shape = slide.shapes.add_shape(
                        MSO_SHAPE.RECTANGLE,
                        Inches(0.7),
                        Inches(0.7 + offset),
                        Inches(2.0),
                        Inches(0.25),
                    )
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = color
                    shape.line.fill.background()
            prs.save(reference)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPTS / "extract_reference_style.py"),
                    str(reference),
                    "--slides",
                    "2",
                    "--out",
                    str(profile),
                ],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            extracted = json.loads(profile.read_text(encoding="utf-8"))
            self.assertEqual(extracted["scope"], "slide_family")
            self.assertEqual(extracted["slide_numbers"], [2])
            self.assertEqual(extracted["slides_analyzed"], 1)
            self.assertEqual(extracted["colors"]["primary"], "B21F2D")

    def test_style_fidelity_qa_passes_for_same_page_family(self):
        with tempfile.TemporaryDirectory() as tmp:
            deck = Path(tmp) / "same-family.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(0.7),
                Inches(0.7),
                Inches(2.0),
                Inches(0.3),
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(0x17, 0x3F, 0x5F)
            shape.line.fill.background()
            textbox = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(5.0), Inches(0.7))
            textbox.text_frame.text = "Architecture-aware mapping"
            run = textbox.text_frame.paragraphs[0].runs[0]
            run.font.name = "Aptos"
            run.font.size = Pt(26)
            run.font.color.rgb = RGBColor(0x17, 0x21, 0x2B)
            prs.save(deck)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPTS / "qa_style_fidelity.py"),
                    str(deck),
                    str(deck),
                    "--generated-slides",
                    "1",
                    "--reference-slides",
                    "1",
                    "--strict",
                ],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            report = json.loads(result.stdout)
            self.assertTrue(report["passed"])
            self.assertEqual(report["metrics"]["palette_coverage"], 1.0)

    def test_semantic_diagram_builds_native_graph_and_traceable_notes(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "semantic-diagram.pptx"
            content_qa = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPTS / "qa_content.py"),
                    str(SEMANTIC_EXAMPLE),
                    "--strict",
                ],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertEqual(content_qa.returncode, 0, content_qa.stdout + content_qa.stderr)
            content_report = json.loads(content_qa.stdout)
            self.assertTrue(content_report["passed"])
            self.assertGreaterEqual(content_report["score"], 80)

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPTS / "build_editable_ppt.py"),
                    str(SEMANTIC_EXAMPLE),
                    "--out",
                    str(out),
                ],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            prs = Presentation(str(out))
            self.assertEqual(len(prs.slides), 3)
            diagram_slide = prs.slides[1]
            names = [str(shape.name) for shape in diagram_slide.shapes]
            self.assertEqual(sum("_NODE_" in name for name in names), 4)
            self.assertEqual(sum("_EDGE_" in name and "_LABEL_" not in name for name in names), 3)
            notes = diagram_slide.notes_slide.notes_text_frame.text
            self.assertIn("example-report.pdf p.2-3", notes)
            self.assertIn("结构化算子", notes)

            qa = subprocess.run(
                [sys.executable, str(SCRIPTS / "qa_editability.py"), str(out), "--strict"],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertEqual(qa.returncode, 0, qa.stdout + qa.stderr)
            qa_report = json.loads(qa.stdout)
            self.assertEqual(qa_report["totals"]["diagram_nodes"], 4)
            self.assertGreaterEqual(qa_report["totals"]["native_connectors"], 3)

    def test_invalid_diagram_edge_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            spec_path = Path(tmp) / "invalid-diagram.json"
            spec_path.write_text(
                json.dumps(
                    {
                        "version": 1,
                        "title": "Invalid diagram",
                        "mode": "native",
                        "slides": [
                            {
                                "type": "diagram",
                                "title": "无效关系会被拒绝",
                                "takeaway": "目标节点必须存在。",
                                "diagram": {
                                    "kind": "decision",
                                    "nodes": [
                                        {"id": "a", "label": "A", "role": "decision"},
                                        {"id": "b", "label": "B", "role": "output"}
                                    ],
                                    "edges": [{"from": "a", "to": "missing"}]
                                }
                            }
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            result = subprocess.run(
                [sys.executable, str(SCRIPTS / "validate_deck_spec.py"), str(spec_path), "--json"],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertNotEqual(result.returncode, 0)
            report = json.loads(result.stdout)
            self.assertFalse(report["valid"])
            self.assertTrue(any("unknown node" in message for message in report["errors"]))

    def test_decomposition_and_composite_diagram_build_as_native_objects(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            spec_path = tmp_path / "composite.json"
            out = tmp_path / "composite.pptx"
            spec_path.write_text(
                json.dumps(
                    {
                        "version": 1,
                        "title": "Parallel and composite semantics",
                        "mode": "native",
                        "style_families": {
                            "technical": {"colors": {"primary": "173F5F", "accent": "B21F2D"}}
                        },
                        "slides": [
                            {
                                "type": "decomposition",
                                "semantic_intent": "decomposition",
                                "style_family": "technical",
                                "title": "三类任务并行形成完整研究闭环",
                                "takeaway": "任务之间相互支撑，但不虚构先后依赖。",
                                "items": [
                                    {"heading": "编程接口", "body": "表达计算语义", "outcome": "统一入口"},
                                    {"heading": "映射优化", "body": "结合硬件约束", "outcome": "高效执行"},
                                    {"heading": "验证评估", "body": "反馈参数调优", "outcome": "闭环验证"},
                                ],
                            },
                            {
                                "type": "composite_diagram",
                                "semantic_intent": "composite",
                                "style_family": "technical",
                                "title": "架构感知映射与端到端验证构成双区协同",
                                "takeaway": "硬件约束进入映射，评估结果反馈多层优化。",
                                "panels": [
                                    {
                                        "heading": "架构感知映射",
                                        "width_ratio": 1.25,
                                        "diagram": {
                                            "kind": "architecture",
                                            "nodes": [
                                                {"id": "front", "label": "Triton 前端", "role": "input"},
                                                {"id": "map", "label": "映射优化", "role": "process"},
                                                {"id": "hw", "label": "硬件约束", "role": "constraint"},
                                            ],
                                            "edges": [
                                                {"from": "front", "to": "map", "type": "data"},
                                                {"from": "hw", "to": "map", "type": "constraint"},
                                            ],
                                        },
                                    },
                                    {
                                        "heading": "验证与调优",
                                        "diagram": {
                                            "kind": "cycle",
                                            "nodes": [
                                                {"id": "run", "label": "运行验证", "role": "process"},
                                                {"id": "eval", "label": "指标评估", "role": "process"},
                                                {"id": "tune", "label": "多层调优", "role": "output"},
                                            ],
                                            "edges": [
                                                {"from": "run", "to": "eval", "type": "flow"},
                                                {"from": "eval", "to": "tune", "type": "flow"},
                                                {"from": "tune", "to": "run", "type": "feedback"},
                                            ],
                                        },
                                    },
                                ],
                            },
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            result = subprocess.run(
                [sys.executable, str(SCRIPTS / "build_editable_ppt.py"), str(spec_path), "--out", str(out)],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            prs = Presentation(str(out))
            self.assertEqual(len(prs.slides), 2)
            first_names = [str(shape.name) for shape in prs.slides[0].shapes]
            second_names = [str(shape.name) for shape in prs.slides[1].shapes]
            self.assertEqual(sum("CPTF_DECOMP_" in name and "_ITEM_" in name for name in first_names), 3)
            self.assertEqual(sum("CPTF_COMPOSITE_" in name and "_NODE_" in name for name in second_names), 6)
            self.assertTrue(any("EDGE_CONSTRAINT" in name for name in second_names))
            self.assertTrue(any("EDGE_FEEDBACK" in name for name in second_names))

    def test_content_qa_flags_false_sequence_duplicates_and_notes_mismatch(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            spec_path = tmp_path / "qa-risks.json"
            out = tmp_path / "qa-risks.pptx"
            repeated = [
                {"heading": "输入规范化"},
                {"heading": "中间表示转换"},
                {"heading": "算子映射"},
            ]
            spec_path.write_text(
                json.dumps(
                    {
                        "version": 1,
                        "slides": [
                            {
                                "type": "process",
                                "semantic_intent": "decomposition",
                                "title": "任务三覆盖三类并行工作",
                                "takeaway": "三类工作共同完成目标。",
                                "steps": repeated,
                                "notes": "方向一聚焦基础平台。",
                            },
                            {
                                "type": "cards",
                                "title": "三项能力共同支撑完整工具链",
                                "takeaway": "能力项与前页重复，需要合并。",
                                "cards": repeated,
                            },
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            result = subprocess.run(
                [sys.executable, str(SCRIPTS / "qa_content.py"), str(spec_path)],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            report = json.loads(result.stdout)
            warnings = "\n".join(report["warnings"])
            self.assertIn("false sequence risk", warnings)
            self.assertIn("structured content duplicates slide 1", warnings)
            self.assertIn("notes reference direction/task 1", warnings)

            build = subprocess.run(
                [sys.executable, str(SCRIPTS / "build_editable_ppt.py"), str(spec_path), "--out", str(out)],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertEqual(build.returncode, 0, build.stdout + build.stderr)
            editability = subprocess.run(
                [sys.executable, str(SCRIPTS / "qa_editability.py"), str(out)],
                check=False,
                text=True,
                capture_output=True,
            )
            self.assertEqual(editability.returncode, 0, editability.stdout + editability.stderr)
            editability_report = json.loads(editability.stdout)
            self.assertTrue(any("notes reference direction/task 1" in item for item in editability_report["warnings"]))


if __name__ == "__main__":
    unittest.main()
