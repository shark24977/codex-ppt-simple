#!/usr/bin/env python3
"""Inspect a PPTX for native editable content and flattened-slide risks."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


_NUMBER_MAP = {"一": "1", "二": "2", "三": "3", "四": "4", "1": "1", "2": "2", "3": "3", "4": "4"}


def _direction_number(value: str) -> str | None:
    match = re.search(r"(?:任务|方向)\s*([一二三四1234])", value or "")
    return _NUMBER_MAP.get(match.group(1)) if match else None


def _slide_title(slide, slide_height: int) -> str:
    try:
        title = slide.shapes.title
        if title is not None and title.text_frame.text.strip():
            return title.text_frame.text.strip()
    except (AttributeError, ValueError):
        pass
    named = [
        shape
        for shape in slide.shapes
        if str(getattr(shape, "name", "")) == "Editable slide title"
        and getattr(shape, "has_text_frame", False)
        and shape.text_frame.text.strip()
    ]
    if named:
        return named[0].text_frame.text.strip()
    candidates = [
        shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
        and shape.text_frame.text.strip()
        and shape.top <= slide_height * 0.3
    ]
    if not candidates:
        return ""
    candidates.sort(key=lambda shape: (shape.top, shape.left))
    return candidates[0].text_frame.text.strip()


def inspect_presentation(path: Path) -> dict:
    prs = Presentation(str(path))
    width = int(prs.slide_width)
    height = int(prs.slide_height)
    report = {
        "pptx": str(path.resolve()),
        "slides": len(prs.slides),
        "totals": {
            "native_text_shapes": 0,
            "native_charts": 0,
            "native_tables": 0,
            "native_auto_shapes": 0,
            "native_connectors": 0,
            "diagram_nodes": 0,
            "native_smartart": 0,
            "pictures": 0,
            "full_slide_pictures": 0,
            "total_shapes": 0,
        },
        "slide_reports": [],
        "warnings": [],
    }

    for slide_index, slide in enumerate(prs.slides, start=1):
        counts = {
            "slide": slide_index,
            "native_text_shapes": 0,
            "native_charts": 0,
            "native_tables": 0,
            "native_auto_shapes": 0,
            "native_connectors": 0,
            "diagram_nodes": 0,
            "native_smartart": 0,
            "pictures": 0,
            "full_slide_pictures": 0,
            "total_shapes": len(slide.shapes),
            "editing_complexity": "normal",
            "speaker_notes": False,
            "notes_alignment": "not_checked",
        }
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                counts["native_text_shapes"] += 1
            if getattr(shape, "has_chart", False):
                counts["native_charts"] += 1
            if getattr(shape, "has_table", False):
                counts["native_tables"] += 1
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                counts["native_auto_shapes"] += 1
            if shape.shape_type == MSO_SHAPE_TYPE.LINE:
                counts["native_connectors"] += 1
            shape_name = str(getattr(shape, "name", ""))
            if shape_name.startswith(("CPTF_DIAGRAM_", "CPTF_SMARTART_", "CPTF_COMPOSITE_")) and "_NODE_" in shape_name:
                counts["diagram_nodes"] += 1
            if shape_name.startswith("CPTF_NATIVE_SMARTART_"):
                counts["native_smartart"] += 1
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                counts["pictures"] += 1
                covers_width = shape.width >= width * 0.92
                covers_height = shape.height >= height * 0.92
                near_origin = shape.left <= width * 0.04 and shape.top <= height * 0.04
                if covers_width and covers_height and near_origin:
                    counts["full_slide_pictures"] += 1
                    report["warnings"].append(
                        f"slide {slide_index} contains a near-full-slide picture; verify that information is not flattened"
                    )
        notes_text = ""
        try:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            counts["speaker_notes"] = bool(notes_text)
        except (AttributeError, ValueError):
            counts["speaker_notes"] = False
        title_text = _slide_title(slide, height)
        title_direction = _direction_number(title_text)
        notes_direction = _direction_number(notes_text)
        if title_direction and notes_direction:
            if title_direction != notes_direction:
                counts["notes_alignment"] = "conflict"
                report["warnings"].append(
                    f"slide {slide_index} notes reference direction/task {notes_direction}, "
                    f"but title references {title_direction}"
                )
            else:
                counts["notes_alignment"] = "ok"
        if counts["native_text_shapes"] == 0:
            report["warnings"].append(f"slide {slide_index} has no native editable text")
        if counts["total_shapes"] > 120 or counts["native_text_shapes"] > 50:
            counts["editing_complexity"] = "high"
            report["warnings"].append(
                f"slide {slide_index} has high editing complexity "
                f"({counts['total_shapes']} shapes, {counts['native_text_shapes']} text shapes); "
                "consider grouping the content into 2-3 editable semantic panels or splitting the slide"
            )
        elif counts["total_shapes"] > 80 or counts["native_text_shapes"] > 35:
            counts["editing_complexity"] = "medium"
            report["warnings"].append(
                f"slide {slide_index} has medium editing complexity "
                f"({counts['total_shapes']} shapes, {counts['native_text_shapes']} text shapes)"
            )
        for key in report["totals"]:
            report["totals"][key] += counts[key]
        report["slide_reports"].append(counts)

    report["editable"] = report["totals"]["full_slide_pictures"] == 0 and all(
        item["native_text_shapes"] > 0 for item in report["slide_reports"]
    )
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", help="PPTX file to inspect")
    parser.add_argument("--json-out", help="Optional JSON report path")
    parser.add_argument("--strict", action="store_true", help="Exit nonzero when flattened-slide risks are found")
    args = parser.parse_args()

    path = Path(args.pptx).expanduser().resolve()
    if not path.is_file():
        print(f"PPTX not found: {path}")
        return 2
    try:
        report = inspect_presentation(path)
    except Exception as exc:
        print(f"Could not inspect PPTX: {exc}")
        return 1

    body = json.dumps(report, ensure_ascii=False, indent=2)
    print(body)
    if args.json_out:
        out = Path(args.json_out).expanduser().resolve()
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(body + "\n", encoding="utf-8")
    return 0 if report["editable"] or not args.strict else 1


if __name__ == "__main__":
    raise SystemExit(main())
