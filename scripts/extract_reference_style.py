#!/usr/bin/env python3
"""Extract editable style tokens from a reference PPTX."""

from __future__ import annotations

import argparse
from collections import Counter
import json
from pathlib import Path
from statistics import median
import xml.etree.ElementTree as ET
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pptx_common import DEFAULT_STYLE, deep_merge, normalize_hex


A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _rgb_from_color(color) -> str | None:
    try:
        rgb = color.rgb
    except (AttributeError, ValueError, TypeError):
        return None
    return str(rgb).upper() if isinstance(rgb, RGBColor) else (str(rgb).upper() if rgb else None)


def _theme_tokens(path: Path) -> tuple[dict[str, str], dict[str, str]]:
    colors: dict[str, str] = {}
    fonts: dict[str, str] = {}
    try:
        with zipfile.ZipFile(path) as archive:
            theme_names = sorted(
                name for name in archive.namelist() if name.startswith("ppt/theme/theme") and name.endswith(".xml")
            )
            if not theme_names:
                return colors, fonts
            root = ET.fromstring(archive.read(theme_names[0]))
    except (OSError, zipfile.BadZipFile, ET.ParseError):
        return colors, fonts

    scheme = root.find(f".//{{{A_NS}}}clrScheme")
    if scheme is not None:
        for item in list(scheme):
            key = item.tag.rsplit("}", 1)[-1]
            color_node = next(iter(item), None)
            if color_node is None:
                continue
            value = color_node.attrib.get("val") or color_node.attrib.get("lastClr")
            if value and len(value) == 6:
                colors[key] = value.upper()

    major = root.find(f".//{{{A_NS}}}majorFont/{{{A_NS}}}latin")
    minor = root.find(f".//{{{A_NS}}}minorFont/{{{A_NS}}}latin")
    if major is not None and major.attrib.get("typeface"):
        fonts["heading"] = major.attrib["typeface"]
    if minor is not None and minor.attrib.get("typeface"):
        fonts["body"] = minor.attrib["typeface"]
    return colors, fonts


def _most_common(counter: Counter[str], fallback: str) -> str:
    return counter.most_common(1)[0][0] if counter else fallback


def _iter_shapes(shapes):
    """Yield top-level shapes and editable children nested inside PowerPoint groups."""

    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def parse_slide_numbers(value: str | None, total: int | None = None) -> list[int] | None:
    """Parse one-based slide selections such as ``2,4-6``."""

    if value is None:
        return None
    numbers: set[int] = set()
    for raw_part in value.split(","):
        part = raw_part.strip()
        if not part:
            continue
        if "-" in part:
            start_text, end_text = part.split("-", 1)
            try:
                start, end = int(start_text), int(end_text)
            except ValueError as exc:
                raise ValueError(f"Invalid slide range: {part!r}") from exc
            if start < 1 or end < start:
                raise ValueError(f"Invalid slide range: {part!r}")
            numbers.update(range(start, end + 1))
        else:
            try:
                number = int(part)
            except ValueError as exc:
                raise ValueError(f"Invalid slide number: {part!r}") from exc
            if number < 1:
                raise ValueError(f"Invalid slide number: {part!r}")
            numbers.add(number)
    if not numbers:
        raise ValueError("Slide selection is empty")
    if total is not None and max(numbers) > total:
        raise ValueError(f"Slide selection exceeds presentation length ({total})")
    return sorted(numbers)


def _colorfulness(value: str) -> tuple[int, int]:
    red, green, blue = (int(value[index : index + 2], 16) for index in (0, 2, 4))
    return max(red, green, blue) - min(red, green, blue), round((red + green + blue) / 3)


def _rank_local_colors(*counters: Counter[str]) -> list[str]:
    combined: Counter[str] = Counter()
    for counter in counters:
        combined.update(counter)
    ranked = []
    for color, _ in combined.most_common():
        spread, brightness = _colorfulness(color)
        if spread >= 24 and 25 <= brightness <= 235:
            ranked.append(color)
    return ranked


def _local_neutral(counter: Counter[str], *, dark: bool, fallback: str) -> str:
    for color, _ in counter.most_common():
        spread, brightness = _colorfulness(color)
        if dark and spread <= 45 and brightness <= 115:
            return color
        if not dark and spread <= 28 and brightness >= 220:
            return color
    return fallback


def extract_style(path: Path, slide_numbers: list[int] | None = None) -> dict:
    prs = Presentation(str(path))
    selected_numbers = slide_numbers or list(range(1, len(prs.slides) + 1))
    if not selected_numbers:
        raise ValueError("At least one slide must be selected")
    if min(selected_numbers) < 1 or max(selected_numbers) > len(prs.slides):
        raise ValueError(f"Slide selection must be within 1-{len(prs.slides)}")
    selected_set = set(selected_numbers)
    local_scope = slide_numbers is not None
    theme_colors, theme_fonts = _theme_tokens(path)
    fill_colors: Counter[str] = Counter()
    line_colors: Counter[str] = Counter()
    background_colors: Counter[str] = Counter()
    font_colors: Counter[str] = Counter()
    heading_fonts: Counter[str] = Counter()
    body_fonts: Counter[str] = Counter()
    heading_sizes: list[float] = []
    body_sizes: list[float] = []

    for slide_number, slide in enumerate(prs.slides, start=1):
        if slide_number not in selected_set:
            continue
        try:
            fill = slide.background.fill
            if fill.type is not None:
                color = _rgb_from_color(fill.fore_color)
                if color:
                    background_colors[color] += 1
        except (AttributeError, ValueError, TypeError):
            pass

        for shape in _iter_shapes(slide.shapes):
            try:
                if shape.fill.type is not None:
                    color = _rgb_from_color(shape.fill.fore_color)
                    if color:
                        fill_colors[color] += 1
            except (AttributeError, ValueError, TypeError):
                pass
            try:
                color = _rgb_from_color(shape.line.color)
                if color:
                    line_colors[color] += 1
            except (AttributeError, ValueError, TypeError):
                pass
            if not getattr(shape, "has_text_frame", False):
                continue
            is_heading_shape = bool(getattr(shape, "is_placeholder", False)) and "title" in str(
                getattr(getattr(shape, "placeholder_format", None), "type", "")
            ).lower()
            for paragraph in shape.text_frame.paragraphs:
                runs = list(paragraph.runs)
                if not runs and paragraph.text:
                    runs = [paragraph]
                for run in runs:
                    font = run.font
                    name = font.name
                    size = font.size.pt if font.size is not None else None
                    color = _rgb_from_color(font.color)
                    heading_like = is_heading_shape or (size is not None and size >= 22)
                    if name:
                        (heading_fonts if heading_like else body_fonts)[name] += 1
                    if size is not None:
                        (heading_sizes if heading_like else body_sizes).append(float(size))
                    if color:
                        font_colors[color] += 1

    fallback = DEFAULT_STYLE
    local_palette = _rank_local_colors(line_colors, fill_colors, font_colors)
    background = _most_common(
        background_colors,
        theme_colors.get("lt1", fallback["colors"]["background"]),
    )
    if local_scope:
        primary = local_palette[0] if local_palette else _most_common(line_colors, fallback["colors"]["primary"])
        accent = next((color for color in local_palette[1:] if color != primary), fallback["colors"]["accent"])
        accent_2 = next(
            (color for color in local_palette[2:] if color not in {primary, accent}),
            fallback["colors"]["accent_2"],
        )
    else:
        primary = theme_colors.get("accent1") or _most_common(fill_colors, fallback["colors"]["primary"])
        accent = theme_colors.get("accent2", fallback["colors"]["accent"])
        accent_2 = theme_colors.get("accent3", fallback["colors"]["accent_2"])
    if local_scope:
        text = _local_neutral(font_colors, dark=True, fallback=fallback["colors"]["text"])
        surface = _local_neutral(fill_colors, dark=False, fallback=fallback["colors"]["surface"])
    else:
        text = theme_colors.get("dk1") or _most_common(font_colors, fallback["colors"]["text"])
        surface = theme_colors.get("lt2", fallback["colors"]["surface"])
    line = _most_common(line_colors, fallback["colors"]["line"])
    muted_text = theme_colors.get("dk2", fallback["colors"]["muted_text"])

    style = deep_merge(
        fallback,
        {
            "colors": {
                "background": normalize_hex(background, fallback["colors"]["background"]),
                "surface": normalize_hex(surface, fallback["colors"]["surface"]),
                "text": normalize_hex(text, fallback["colors"]["text"]),
                "muted_text": normalize_hex(muted_text, fallback["colors"]["muted_text"]),
                "primary": normalize_hex(primary, fallback["colors"]["primary"]),
                "accent": normalize_hex(accent, fallback["colors"]["accent"]),
                "accent_2": normalize_hex(accent_2, fallback["colors"]["accent_2"]),
                "line": normalize_hex(line, fallback["colors"]["line"]),
            },
            "fonts": {
                "heading": _most_common(heading_fonts, theme_fonts.get("heading", fallback["fonts"]["heading"])),
                "body": _most_common(body_fonts, theme_fonts.get("body", fallback["fonts"]["body"])),
            },
            "sizes": {
                "title": round(median(heading_sizes), 1) if heading_sizes else fallback["sizes"]["title"],
                "body": round(median(body_sizes), 1) if body_sizes else fallback["sizes"]["body"],
            },
        },
    )
    return {
        "profile_version": 1,
        "source": str(path.resolve()),
        "scope": "slide_family" if local_scope else "whole_deck",
        "slide_numbers": selected_numbers,
        "slides_analyzed": len(selected_numbers),
        "slide_size_inches": {
            "width": round(prs.slide_width / 914400, 4),
            "height": round(prs.slide_height / 914400, 4),
        },
        **style,
        "evidence": {
            "theme_colors": theme_colors,
            "common_fill_colors": [color for color, _ in fill_colors.most_common(8)],
            "common_line_colors": [color for color, _ in line_colors.most_common(8)],
            "common_font_colors": [color for color, _ in font_colors.most_common(8)],
            "local_palette_candidates": local_palette[:8],
            "common_heading_fonts": [font for font, _ in heading_fonts.most_common(5)],
            "common_body_fonts": [font for font, _ in body_fonts.most_common(5)],
        },
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("reference_pptx", help="Editable reference PPTX")
    parser.add_argument("--out", required=True, help="Output style profile JSON")
    parser.add_argument(
        "--slides",
        help="One-based representative slide family, for example 10-13,16; local colors override theme accents",
    )
    args = parser.parse_args()

    reference = Path(args.reference_pptx).expanduser().resolve()
    if not reference.is_file():
        print(f"Reference PPTX not found: {reference}")
        return 2
    try:
        slide_numbers = parse_slide_numbers(args.slides)
        profile = extract_style(reference, slide_numbers=slide_numbers)
    except Exception as exc:
        print(f"Could not extract style: {exc}")
        return 1
    out = Path(args.out).expanduser().resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(profile, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote style profile: {out}")
    print(f"slides_analyzed={profile['slides_analyzed']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
