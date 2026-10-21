#!/usr/bin/env python3
"""Inspect a generated PPTX for font-size, density, and emphasis hierarchy."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import re
from typing import Any, Iterable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pptx_common import (
    DEFAULT_DENSITY_POLICY,
    DEFAULT_EMPHASIS_POLICY,
    DEFAULT_READABILITY_POLICY,
    load_json,
    merged_policy,
)


CONTENT_EXEMPT_TYPES = {"cover", "section", "closing"}


def _compact_length(value: str) -> int:
    return len(re.sub(r"\s+", "", value or ""))


def _shape_role(shape, *, is_title: bool = False) -> str:
    name = str(getattr(shape, "name", "")).lower()
    if "panel" in name and "title" in name:
        return "module"
    if is_title or ("title" in name and "subtitle" not in name):
        return "title"
    if "panel takeaway" in name:
        return "label"
    if "takeaway" in name or "subtitle" in name:
        return "takeaway"
    if "heading" in name:
        return "module"
    if any(token in name for token in ("footer", "number", "kicker", "label", "outcome", "contact")):
        return "label"
    if name.startswith(("cptf_diagram_", "cptf_smartart_", "cptf_composite_")):
        return "diagram"
    return "body"


def _iter_shapes(shapes) -> Iterable[Any]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _paragraph_runs(paragraph) -> Iterable[tuple[str, float | None, bool]]:
    if paragraph.runs:
        for run in paragraph.runs:
            text = run.text or ""
            size = run.font.size or paragraph.font.size
            bold = run.font.bold if run.font.bold is not None else paragraph.font.bold
            yield text, (float(size.pt) if size is not None else None), bool(bold)
    elif paragraph.text:
        size = paragraph.font.size
        yield paragraph.text, (float(size.pt) if size is not None else None), bool(paragraph.font.bold)


def _minimum_for_role(policy: dict[str, Any], role: str) -> float:
    key = {
        "title": "min_title_font",
        "takeaway": "min_takeaway_font",
        "module": "min_module_font",
        "body": "min_body_font",
        "diagram": "min_diagram_font",
        "label": "min_label_font",
    }.get(role, "min_body_font")
    return float(policy[key])


def inspect_presentation(path: Path, spec: dict[str, Any] | None = None) -> dict[str, Any]:
    spec = spec or {}
    readability = merged_policy(spec, "readability_policy", DEFAULT_READABILITY_POLICY)
    emphasis = merged_policy(spec, "emphasis_policy", DEFAULT_EMPHASIS_POLICY)
    density = merged_policy(spec, "density_policy", DEFAULT_DENSITY_POLICY)
    prs = Presentation(str(path))
    slide_specs = spec.get("slides") if isinstance(spec.get("slides"), list) else []
    reports: list[dict[str, Any]] = []
    warnings: list[str] = []
    violation_count = 0
    total_undersized = 0

    for slide_index, slide in enumerate(prs.slides, start=1):
        slide_spec = slide_specs[slide_index - 1] if slide_index <= len(slide_specs) else {}
        slide_type = str(slide_spec.get("type", "content")) if isinstance(slide_spec, dict) else "content"
        title_shape = None
        try:
            title_shape = slide.shapes.title
        except (AttributeError, ValueError):
            pass
        visible_characters = 0
        emphasis_characters = 0
        emphasis_scope_characters = 0
        undersized: list[dict[str, Any]] = []
        unknown_font_runs = 0

        for shape in _iter_shapes(slide.shapes):
            role = _shape_role(shape, is_title=shape is title_shape)
            frames: list[tuple[Any, str]] = []
            if getattr(shape, "has_text_frame", False):
                frames.append((shape.text_frame, role))
            if getattr(shape, "has_table", False):
                for row_index, row in enumerate(shape.table.rows):
                    for cell in row.cells:
                        frames.append((cell.text_frame, "module" if row_index == 0 else "body"))

            for frame, frame_role in frames:
                minimum = _minimum_for_role(readability, frame_role)
                for paragraph in frame.paragraphs:
                    for text, font_size, bold in _paragraph_runs(paragraph):
                        characters = _compact_length(text)
                        if not characters:
                            continue
                        if frame_role != "label":
                            visible_characters += characters
                        if frame_role not in {"title", "label"}:
                            emphasis_scope_characters += characters
                            if bold:
                                emphasis_characters += characters
                        if font_size is None:
                            unknown_font_runs += 1
                        elif font_size + 0.01 < minimum:
                            undersized.append(
                                {
                                    "shape": str(getattr(shape, "name", "")),
                                    "role": frame_role,
                                    "font_size": round(font_size, 2),
                                    "minimum": minimum,
                                    "text": text.strip()[:80],
                                }
                            )

        exempt = slide_type in CONTENT_EXEMPT_TYPES
        bold_ratio = (
            emphasis_characters / emphasis_scope_characters if emphasis_scope_characters else 0.0
        )
        issues: list[str] = []
        if undersized:
            issues.append(f"{len(undersized)} text run(s) are below their role-specific font minimum")
        if not exempt and visible_characters > int(density["max_visible_characters"]):
            issues.append(
                f"{visible_characters} visible characters exceed density limit "
                f"{density['max_visible_characters']}; revise or split instead of shrinking fonts"
            )
        if (
            not exempt
            and bool(emphasis.get("enabled", True))
            and emphasis_scope_characters >= 60
            and bold_ratio + 0.001 < float(emphasis["min_bold_ratio"])
        ):
            issues.append(
                f"bold emphasis ratio {bold_ratio:.1%} is below {float(emphasis['min_bold_ratio']):.1%}"
            )
        if (
            not exempt
            and emphasis_scope_characters >= 180
            and bold_ratio > float(emphasis["max_bold_ratio"])
        ):
            warnings.append(
                f"slide {slide_index}: bold emphasis ratio {bold_ratio:.1%} exceeds "
                f"{float(emphasis['max_bold_ratio']):.1%}; avoid making most content equally prominent"
            )

        if issues:
            violation_count += len(issues)
            warnings.extend(f"slide {slide_index}: {issue}" for issue in issues)
        total_undersized += len(undersized)
        reports.append(
            {
                "slide": slide_index,
                "type": slide_type,
                "visible_characters": visible_characters,
                "emphasis_scope_characters": emphasis_scope_characters,
                "bold_characters": emphasis_characters,
                "bold_ratio": round(bold_ratio, 4),
                "undersized_runs": undersized,
                "unknown_font_runs": unknown_font_runs,
                "issues": issues,
            }
        )

    return {
        "pptx": str(path.resolve()),
        "slides": len(prs.slides),
        "passed": violation_count == 0,
        "policies": {
            "readability": readability,
            "emphasis": emphasis,
            "density": density,
        },
        "totals": {
            "violations": violation_count,
            "undersized_runs": total_undersized,
        },
        "warnings": warnings,
        "slide_reports": reports,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", help="PPTX file to inspect")
    parser.add_argument("--spec", help="Optional deck_spec.json supplying readability policies and slide types")
    parser.add_argument("--json-out", help="Optional JSON report path")
    parser.add_argument("--strict", action="store_true", help="Exit nonzero when readability violations are found")
    args = parser.parse_args()

    path = Path(args.pptx).expanduser().resolve()
    if not path.is_file():
        print(f"PPTX not found: {path}")
        return 2
    spec: dict[str, Any] | None = None
    if args.spec:
        try:
            spec = load_json(Path(args.spec).expanduser().resolve())
        except ValueError as exc:
            print(str(exc))
            return 2
    try:
        report = inspect_presentation(path, spec)
    except Exception as exc:
        print(f"Could not inspect PPTX readability: {exc}")
        return 1

    body = json.dumps(report, ensure_ascii=False, indent=2)
    print(body)
    if args.json_out:
        out = Path(args.json_out).expanduser().resolve()
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(body + "\n", encoding="utf-8")
    return 0 if report["passed"] or not args.strict else 1


if __name__ == "__main__":
    raise SystemExit(main())
