#!/usr/bin/env python3
"""Build a native editable PowerPoint deck from deck_spec.json."""

from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path
from typing import Any, Dict, Iterable, List

from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from diagram_layout import layout_diagram
from extract_reference_style import extract_style
from pptx_common import (
    DEFAULT_EMPHASIS_POLICY,
    DEFAULT_READABILITY_POLICY,
    load_json,
    merged_policy,
    merged_style,
    resolve_asset_path,
    validate_spec,
)


EMU_PER_INCH = 914400


def _remove_all_slides(prs: Presentation) -> None:
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        slide_ids.remove(slide_id)


def _blank_layout(prs: Presentation):
    named = [layout for layout in prs.slide_layouts if str(layout.name).strip().lower() in {"blank", "空白"}]
    if named:
        return named[0]
    return min(prs.slide_layouts, key=lambda layout: len(layout.placeholders))


def _inches(value: int) -> float:
    return value / EMU_PER_INCH


class DeckBuilder:
    def __init__(self, prs: Presentation, spec: Dict[str, Any], style: Dict[str, Any], spec_path: Path):
        self.prs = prs
        self.spec = spec
        self.base_style = style
        self.style = style
        self.spec_path = spec_path
        self.layout = _blank_layout(prs)
        self.width = _inches(prs.slide_width)
        self.height = _inches(prs.slide_height)
        self.colors = style["colors"]
        self.fonts = style["fonts"]
        self.sizes = style["sizes"]
        self.geo = style["geometry"]
        self.readability = merged_policy(spec, "readability_policy", DEFAULT_READABILITY_POLICY)
        self.emphasis = merged_policy(spec, "emphasis_policy", DEFAULT_EMPHASIS_POLICY)

    def _set_style(self, style: Dict[str, Any]) -> None:
        self.style = style
        self.colors = style["colors"]
        self.fonts = style["fonts"]
        self.sizes = style["sizes"]
        self.geo = style["geometry"]

    def _style_for_slide(self, data: Dict[str, Any]) -> Dict[str, Any]:
        family_name = data.get("style_family")
        families = self.spec.get("style_families", {})
        family = families.get(str(family_name), {}) if isinstance(families, dict) and family_name else {}
        style = merged_style(self.base_style, family)
        return merged_style(style, data.get("style") if isinstance(data.get("style"), dict) else {})

    @staticmethod
    def rgb(value: str) -> RGBColor:
        value = value.lstrip("#")
        return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))

    @staticmethod
    def _text_role(name: str | None) -> str:
        value = str(name or "").lower()
        if "panel takeaway" in value:
            return "label"
        if "panel" in value and "title" in value:
            return "module"
        if "title" in value and "subtitle" not in value:
            return "title"
        if "takeaway" in value or "subtitle" in value:
            return "takeaway"
        if "heading" in value:
            return "module"
        if any(token in value for token in ("footer", "number", "kicker", "label", "outcome", "contact")):
            return "label"
        if "diagram" in value:
            return "diagram"
        return "body"

    def _font_size(self, proposed: float, role: str | None = None, name: str | None = None) -> float:
        resolved_role = role or self._text_role(name)
        minimums = {
            "title": float(self.readability["min_title_font"]),
            "takeaway": float(self.readability["min_takeaway_font"]),
            "module": float(self.readability["min_module_font"]),
            "body": float(self.readability["min_body_font"]),
            "diagram": float(self.readability["min_diagram_font"]),
            "label": float(self.readability["min_label_font"]),
        }
        return max(float(proposed), minimums.get(resolved_role, minimums["body"]))

    def _rich_segments(self, value: Any, *, base_bold: bool, auto_label: bool) -> List[tuple[str, bool]]:
        text = str(value or "")
        if base_bold:
            return [(text.replace("**", ""), True)]
        segments: List[tuple[str, bool]] = []
        cursor = 0
        for match in re.finditer(r"\*\*(.+?)\*\*", text, flags=re.DOTALL):
            if match.start() > cursor:
                segments.append((text[cursor : match.start()], False))
            segments.append((match.group(1), True))
            cursor = match.end()
        if cursor:
            if cursor < len(text):
                segments.append((text[cursor:], False))
            return [(part, bold) for part, bold in segments if part]
        if auto_label and bool(self.emphasis.get("auto_emphasize_leading_labels", True)):
            match = re.match(r"^(\s*[•·▪-]?\s*)([^：:\n]{1,18})([：:])(.*)$", text, flags=re.DOTALL)
            if match:
                prefix, label, colon, remainder = match.groups()
                return [
                    (prefix, False),
                    (label + colon, True),
                    (remainder, False),
                ]
        return [(text.replace("**", ""), False)]

    def _format_paragraph(
        self,
        paragraph,
        text: Any,
        *,
        size: float,
        color: str,
        font: str,
        bold: bool,
        auto_label: bool = False,
    ) -> None:
        paragraph.text = ""
        for value, run_bold in self._rich_segments(text, base_bold=bold, auto_label=auto_label):
            run = paragraph.add_run()
            run.text = value
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = run_bold
            run.font.color.rgb = self.rgb(color)

    def add_rect(
        self,
        slide,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        fill: str,
        line: str | None = None,
        rounded: bool = False,
        name: str | None = None,
    ):
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.rgb(fill)
        if line:
            shape.line.color.rgb = self.rgb(line)
            shape.line.width = Pt(0.8)
        else:
            shape.line.fill.background()
        if name:
            shape.name = name
        return shape

    def add_native_shape(
        self,
        slide,
        shape_type,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        fill: str,
        line: str | None = None,
        name: str | None = None,
    ):
        shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.rgb(fill)
        if line:
            shape.line.color.rgb = self.rgb(line)
            shape.line.width = Pt(1.15)
        else:
            shape.line.fill.background()
        if name:
            shape.name = name
        return shape

    @staticmethod
    def _set_end_arrow(connector) -> None:
        line = connector.line._get_or_add_ln()
        existing = line.find(qn("a:tailEnd"))
        if existing is not None:
            line.remove(existing)
        arrow = OxmlElement("a:tailEnd")
        arrow.set("type", "triangle")
        arrow.set("w", "sm")
        arrow.set("len", "sm")
        line.append(arrow)

    @staticmethod
    def _set_dash(connector, dashed: bool) -> None:
        if not dashed:
            return
        line = connector.line._get_or_add_ln()
        existing = line.find(qn("a:prstDash"))
        if existing is not None:
            line.remove(existing)
        dash = OxmlElement("a:prstDash")
        dash.set("val", "dash")
        line.append(dash)

    def add_connector(
        self,
        slide,
        x1: float,
        y1: float,
        x2: float,
        y2: float,
        *,
        elbow: bool = False,
        color: str | None = None,
        dashed: bool = False,
        name: str | None = None,
    ):
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT,
            Inches(x1),
            Inches(y1),
            Inches(x2),
            Inches(y2),
        )
        connector.line.color.rgb = self.rgb(color or self.colors["primary"])
        connector.line.width = Pt(1.45)
        self._set_end_arrow(connector)
        self._set_dash(connector, dashed)
        if name:
            connector.name = name
        return connector

    def add_text(
        self,
        slide,
        text: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        size: float,
        color: str | None = None,
        bold: bool = False,
        font: str | None = None,
        align=PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.TOP,
        name: str | None = None,
        margin: float = 0.0,
        role: str | None = None,
    ):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            shape.name = name
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.vertical_anchor = valign
        frame.margin_left = Inches(margin)
        frame.margin_right = Inches(margin)
        frame.margin_top = Inches(margin)
        frame.margin_bottom = Inches(margin)
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        resolved_role = role or self._text_role(name)
        resolved_size = self._font_size(size, resolved_role, name)
        self._format_paragraph(
            paragraph,
            text,
            size=resolved_size,
            color=color or self.colors["text"],
            font=font or self.fonts["body"],
            bold=bold,
            auto_label=resolved_role == "body" and bool(self.emphasis.get("enabled", True)),
        )
        return shape

    def add_bullets(
        self,
        slide,
        bullets: Iterable[Any],
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        size: float | None = None,
        color: str | None = None,
        name: str = "Editable bullets",
    ):
        items = [str(item).strip() for item in bullets if str(item).strip()]
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = name
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = Inches(0.04)
        frame.margin_right = Inches(0.04)
        frame.margin_top = Inches(0.02)
        frame.margin_bottom = Inches(0.02)
        for index, item in enumerate(items):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            self._format_paragraph(
                paragraph,
                f"•  {item}",
                size=self._font_size(size or self.sizes["body"], "body", name),
                color=color or self.colors["text"],
                font=self.fonts["body"],
                bold=False,
                auto_label=bool(self.emphasis.get("enabled", True)),
            )
            paragraph.space_after = Pt(7)
            paragraph.line_spacing = 1.1
        return shape

    def set_background(self, slide, color: str | None = None) -> None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self.rgb(color or self.colors["background"])

    def add_header(self, slide, data: Dict[str, Any]) -> float:
        header = self.style.get("header", {})
        mx = float(header.get("title_x", self.geo["margin_x"]))
        mt = float(header.get("title_y", self.geo["margin_top"]))
        title_height = float(header.get("title_height", self.geo["title_height"]))
        title_color_role = str(header.get("title_color_role", "text"))
        title_color = self.colors.get(title_color_role, self.colors["text"])
        if data.get("kicker"):
            self.add_text(
                slide,
                data["kicker"],
                mx,
                mt - 0.05,
                self.width - 2 * mx,
                0.25,
                size=self.sizes["small"],
                color=self.colors["accent"],
                bold=True,
                name="Editable kicker",
            )
            title_y = mt + 0.2
        else:
            title_y = mt
        self.add_text(
            slide,
            data.get("title", ""),
            mx,
            title_y,
            self.width - mx - float(header.get("right_margin", mx)),
            title_height,
            size=self.sizes["title"],
            color=title_color,
            bold=True,
            font=self.fonts["heading"],
            name="Editable slide title",
        )
        if data.get("subtitle"):
            subtitle_y = title_y + title_height + float(header.get("subtitle_gap", 0.08))
            self.add_text(
                slide,
                data["subtitle"],
                mx,
                subtitle_y,
                self.width - mx - float(header.get("right_margin", mx)),
                float(header.get("subtitle_height", 0.4)),
                size=self.sizes["subtitle"],
                color=self.colors["muted_text"],
                name="Editable slide subtitle",
            )
            return subtitle_y + float(header.get("subtitle_height", 0.4)) + float(header.get("content_gap", 0.12))

        if header.get("full_width_rule"):
            rule_y = float(header.get("rule_y", title_y + title_height + 0.05))
            self.add_rect(
                slide,
                float(header.get("rule_x", 0.0)),
                rule_y,
                float(header.get("rule_width", self.width)),
                float(header.get("rule_height", 0.045)),
                fill=self.colors.get(str(header.get("rule_color_role", "primary")), self.colors["primary"]),
                name="Editable full-width title rule",
            )
            return rule_y + float(header.get("content_gap", 0.28))

        self.add_rect(
            slide,
            mx,
            title_y + title_height + 0.05,
            0.72,
            0.055,
            fill=self.colors["accent"],
            name="Editable title accent",
        )
        return title_y + title_height + 0.25

    def add_footer(self, slide, data: Dict[str, Any], slide_number: int) -> None:
        show_number = data.get("show_slide_number", self.spec.get("show_slide_numbers", True))
        y = self.height - 0.33
        if data.get("footer"):
            self.add_text(
                slide,
                data["footer"],
                float(self.geo["margin_x"]),
                y,
                self.width - 2.0,
                0.18,
                size=self.sizes["small"],
                color=self.colors["muted_text"],
                name="Editable footer",
            )
        if show_number:
            self.add_text(
                slide,
                slide_number,
                self.width - float(self.geo["margin_x"]) - 0.35,
                y,
                0.35,
                0.18,
                size=self.sizes["small"],
                color=self.colors["muted_text"],
                align=PP_ALIGN.RIGHT,
                name="Editable slide number",
            )

    def add_takeaway(self, slide, text: Any) -> None:
        if not str(text or "").strip():
            return
        mx = float(self.geo["margin_x"])
        self.add_rect(
            slide,
            mx,
            self.height - 0.92,
            self.width - 2 * mx,
            0.45,
            fill=self.colors["surface"],
            line=self.colors["line"],
            rounded=True,
            name="Editable takeaway container",
        )
        self.add_text(
            slide,
            text,
            mx + 0.18,
            self.height - 0.83,
            self.width - 2 * mx - 0.36,
            0.25,
            size=max(11, self.sizes["body"] - 1),
            color=self.colors["primary"],
            bold=True,
            name="Editable takeaway",
        )

    def add_image_contain(self, slide, image: Dict[str, Any] | str, x: float, y: float, w: float, h: float):
        value = {"path": image} if isinstance(image, str) else image
        raw_path = value.get("path") if isinstance(value, dict) else None
        if not raw_path:
            raise ValueError("image.path is required")
        path = resolve_asset_path(self.spec_path, str(raw_path))
        if not path.is_file():
            raise ValueError(f"Image asset not found: {path}")
        with Image.open(path) as raster:
            pixel_w, pixel_h = raster.size
        scale = min(w / pixel_w, h / pixel_h)
        placed_w = pixel_w * scale
        placed_h = pixel_h * scale
        left = x + (w - placed_w) / 2
        top = y + (h - placed_h) / 2
        picture = slide.shapes.add_picture(
            str(path),
            Inches(left),
            Inches(top),
            width=Inches(placed_w),
            height=Inches(placed_h),
        )
        picture.name = str(value.get("alt") or "Replaceable illustration")
        return picture

    def _new_slide(self, background: str | None = None):
        slide = self.prs.slides.add_slide(self.layout)
        self.set_background(slide, background)
        return slide

    def _add_notes(self, slide, data: Dict[str, Any]) -> None:
        sections: List[str] = []
        notes = str(data.get("notes") or "").strip()
        if notes:
            sections.append(notes)
        profile = self.spec.get("content_profile", {})
        include_evidence = not isinstance(profile, dict) or profile.get("notes_include_evidence", True)
        if include_evidence:
            raw_evidence = data.get("evidence", [])
            evidence_values = raw_evidence if isinstance(raw_evidence, list) else [raw_evidence]
            evidence = [str(item).strip() for item in evidence_values if str(item).strip()]
            if evidence:
                sections.append("证据：\n" + "\n".join(f"- {item}" for item in evidence))
        raw_refs = data.get("source_refs", [])
        ref_values = raw_refs if isinstance(raw_refs, list) else [raw_refs]
        source_refs = [str(item).strip() for item in ref_values if str(item).strip()]
        if source_refs:
            sections.append("来源：" + "；".join(source_refs))
        if not sections:
            return
        frame = slide.notes_slide.notes_text_frame
        frame.clear()
        frame.text = "\n\n".join(sections)

    def render_cover(self, data: Dict[str, Any], slide_number: int):
        slide = self._new_slide(self.colors["primary"])
        mx = float(self.geo["margin_x"])
        text_w = self.width * (0.55 if data.get("image") else 0.82)
        if data.get("kicker"):
            self.add_text(
                slide,
                data["kicker"],
                mx,
                0.95,
                text_w,
                0.3,
                size=self.sizes["small"] + 1,
                color=self.colors["accent_2"],
                bold=True,
                name="Editable cover kicker",
            )
        self.add_text(
            slide,
            data.get("title", ""),
            mx,
            1.45,
            text_w,
            1.75,
            size=self.sizes["cover_title"],
            color="FFFFFF",
            bold=True,
            font=self.fonts["heading"],
            valign=MSO_ANCHOR.MIDDLE,
            name="Editable cover title",
        )
        self.add_text(
            slide,
            data.get("subtitle", ""),
            mx,
            3.45,
            text_w,
            0.9,
            size=self.sizes["subtitle"],
            color="E8EEF3",
            name="Editable cover subtitle",
        )
        self.add_rect(slide, mx, 4.65, 1.1, 0.08, fill=self.colors["accent"], name="Editable cover accent")
        if data.get("image"):
            self.add_rect(
                slide,
                self.width * 0.62,
                0.85,
                self.width * 0.31,
                self.height - 1.7,
                fill=self.colors["surface"],
                rounded=True,
                name="Editable illustration container",
            )
            self.add_image_contain(
                slide,
                data["image"],
                self.width * 0.64,
                1.0,
                self.width * 0.27,
                self.height - 2.0,
            )
        if data.get("show_slide_number"):
            self.add_footer(slide, data, slide_number)
        return slide

    def render_section(self, data: Dict[str, Any], slide_number: int):
        slide = self._new_slide(self.colors["primary"])
        mx = float(self.geo["margin_x"])
        if data.get("section_number") is not None:
            self.add_text(
                slide,
                str(data["section_number"]),
                mx,
                1.2,
                1.0,
                0.8,
                size=36,
                color=self.colors["accent_2"],
                bold=True,
                name="Editable section number",
            )
        self.add_text(
            slide,
            data.get("title", ""),
            mx,
            2.15,
            self.width - 2 * mx,
            1.1,
            size=self.sizes["cover_title"],
            color="FFFFFF",
            bold=True,
            font=self.fonts["heading"],
            name="Editable section title",
        )
        self.add_text(
            slide,
            data.get("subtitle", ""),
            mx,
            3.45,
            self.width - 2 * mx,
            0.7,
            size=self.sizes["subtitle"],
            color="E8EEF3",
            name="Editable section subtitle",
        )
        self.add_rect(slide, mx, 4.55, 1.1, 0.08, fill=self.colors["accent"], name="Editable section accent")
        self.add_footer(slide, data, slide_number)
        return slide

    def render_title_body(self, data: Dict[str, Any], slide_number: int):
        slide = self._new_slide()
        content_y = self.add_header(slide, data)
        mx = float(self.geo["margin_x"])
        bottom = self.height - (1.05 if data.get("takeaway") else 0.55)
        content_h = bottom - content_y
        if data.get("image"):
            left_w = (self.width - 2 * mx) * 0.55
            self.add_bullets(slide, data.get("bullets", []), mx, content_y + 0.12, left_w, content_h - 0.15)
            image_x = mx + left_w + float(self.geo["gutter"])
            image_w = self.width - mx - image_x
            self.add_rect(
                slide,
                image_x,
                content_y,
                image_w,
                content_h,
                fill=self.colors["surface"],
                line=self.colors["line"],
                rounded=True,
                name="Editable illustration container",
            )
            self.add_image_contain(slide, data["image"], image_x + 0.15, content_y + 0.15, image_w - 0.3, content_h - 0.3)
        else:
            self.add_bullets(slide, data.get("bullets", []), mx, content_y + 0.12, self.width - 2 * mx, content_h - 0.15)
        self.add_takeaway(slide, data.get("takeaway"))
        self.add_footer(slide, data, slide_number)
        return slide

    def _render_columns(self, data: Dict[str, Any], slide_number: int, comparison: bool):
        slide = self._new_slide()
        content_y = self.add_header(slide, data)
        mx = float(self.geo["margin_x"])
        gutter = float(self.geo["gutter"])
        w = (self.width - 2 * mx - gutter) / 2
        h = self.height - content_y - (1.05 if data.get("takeaway") else 0.55)
        for index, (side, x) in enumerate(((data.get("left", {}), mx), (data.get("right", {}), mx + w + gutter))):
            accent = side.get("accent") or (self.colors["accent"] if index == 0 else self.colors["accent_2"])
            self.add_rect(slide, x, content_y, w, h, fill=self.colors["surface"], line=self.colors["line"], rounded=True)
            self.add_rect(slide, x, content_y, 0.08, h, fill=accent, name="Editable column accent")
            self.add_text(
                slide,
                side.get("heading", ""),
                x + 0.26,
                content_y + 0.22,
                w - 0.5,
                0.42,
                size=self.sizes["subtitle"],
                color=self.colors["primary"],
                bold=True,
                font=self.fonts["heading"],
                name="Editable column heading",
            )
            self.add_bullets(
                slide,
                side.get("bullets", []),
                x + 0.26,
                content_y + 0.82,
                w - 0.52,
                h - 1.02,
                size=self.sizes["body"],
                name="Editable comparison bullets" if comparison else "Editable column bullets",
            )
        self.add_takeaway(slide, data.get("takeaway"))
        self.add_footer(slide, data, slide_number)
        return slide

    def render_cards(self, data: Dict[str, Any], slide_number: int):
        slide = self._new_slide()
        content_y = self.add_header(slide, data)
        cards = data.get("cards", [])
        mx = float(self.geo["margin_x"])
        gutter = float(self.geo["gutter"])
        cols = 3 if len(cards) in {3, 5, 6} else 2
        rows = math.ceil(len(cards) / cols)
        available_h = self.height - content_y - (1.05 if data.get("takeaway") else 0.55)
        card_w = (self.width - 2 * mx - gutter * (cols - 1)) / cols
        card_h = (available_h - gutter * (rows - 1)) / rows
        for index, card in enumerate(cards):
            row, col = divmod(index, cols)
            x = mx + col * (card_w + gutter)
            y = content_y + row * (card_h + gutter)
            accent = card.get("accent") or [self.colors["accent"], self.colors["accent_2"], self.colors["primary"]][index % 3]
            self.add_rect(slide, x, y, card_w, card_h, fill=self.colors["surface"], line=self.colors["line"], rounded=True)
            self.add_rect(slide, x + 0.18, y + 0.2, 0.48, 0.055, fill=accent)
            cursor_y = y + 0.38
            if card.get("value"):
                self.add_text(
                    slide,
                    card["value"],
                    x + 0.18,
                    cursor_y,
                    card_w - 0.36,
                    0.45,
                    size=self.sizes["card_value"],
                    color=accent,
                    bold=True,
                    font=self.fonts["heading"],
                    name="Editable card value",
                )
                cursor_y += 0.58
            self.add_text(
                slide,
                card.get("heading", ""),
                x + 0.18,
                cursor_y,
                card_w - 0.36,
                0.42,
                size=max(14, self.sizes["body"] + 1),
                color=self.colors["text"],
                bold=True,
                name="Editable card heading",
            )
            self.add_text(
                slide,
                card.get("body", ""),
                x + 0.18,
                cursor_y + 0.52,
                card_w - 0.36,
                max(0.45, card_h - (cursor_y - y) - 0.7),
                size=max(10.5, self.sizes["body"] - 1),
                color=self.colors["muted_text"],
                name="Editable card body",
            )
        self.add_takeaway(slide, data.get("takeaway"))
        self.add_footer(slide, data, slide_number)
        return slide

    def render_process(self, data: Dict[str, Any], slide_number: int):
        slide = self._new_slide()
        content_y = self.add_header(slide, data)
        steps = data.get("steps", [])
        mx = float(self.geo["margin_x"])
        gap = 0.18
        arrow_w = 0.28
        total_arrows = arrow_w * max(0, len(steps) - 1)
        card_w = (self.width - 2 * mx - gap * max(0, len(steps) - 1) - total_arrows) / len(steps)
        card_h = min(3.35, self.height - content_y - 1.25)
        y = content_y + 0.55
        compact = len(steps) >= 5
        heading_size = 13 if compact else max(13, self.sizes["body"])
        body_size = 10.5 if compact else max(10.5, self.sizes["body"] - 2)
        for index, step in enumerate(steps):
            x = mx + index * (card_w + gap + arrow_w)
            self.add_rect(slide, x, y, card_w, card_h, fill=self.colors["surface"], line=self.colors["line"], rounded=True)
            self.add_rect(slide, x + 0.18, y + 0.22, 0.45, 0.45, fill=self.colors["primary"], rounded=True)
            self.add_text(
                slide,
                index + 1,
                x + 0.18,
                y + 0.27,
                0.45,
                0.28,
                size=12,
                color="FFFFFF",
                bold=True,
                align=PP_ALIGN.CENTER,
                name="Editable process number",
            )
            self.add_text(
                slide,
                step.get("heading", ""),
                x + 0.18,
                y + 0.88,
                card_w - 0.36,
                0.55,
                size=heading_size,
                color=self.colors["text"],
                bold=True,
                name="Editable process heading",
            )
            self.add_text(
                slide,
                step.get("body", ""),
                x + 0.18,
                y + 1.55,
                card_w - 0.36,
                card_h - 1.8,
                size=body_size,
                color=self.colors["muted_text"],
                name="Editable process body",
            )
            if index < len(steps) - 1:
                arrow_x = x + card_w + gap / 2
                arrow = slide.shapes.add_shape(
                    MSO_SHAPE.CHEVRON,
                    Inches(arrow_x),
                    Inches(y + card_h / 2 - 0.18),
                    Inches(arrow_w),
                    Inches(0.36),
                )
                arrow.fill.solid()
                arrow.fill.fore_color.rgb = self.rgb(self.colors["accent"])
                arrow.line.fill.background()
                arrow.name = "Editable process arrow"
        self.add_takeaway(slide, data.get("takeaway"))
        self.add_footer(slide, data, slide_number)
        return slide

    def render_decomposition(self, data: Dict[str, Any], slide_number: int):
        """Render parallel work packages without implying sequence or precedence."""

        slide = self._new_slide()
        content_y = self.add_header(slide, data)
        items = data.get("items", [])
        mx = float(self.geo["margin_x"])
        gutter = float(self.geo["gutter"])
        cols = 3 if len(items) in {3, 5, 6} else 2
        rows = math.ceil(len(items) / cols)
        bottom = self.height - (1.05 if data.get("takeaway") else 0.55)
        available_h = max(1.0, bottom - content_y)
        card_w = (self.width - 2 * mx - gutter * (cols - 1)) / cols
        card_h = (available_h - gutter * (rows - 1)) / rows
        palette = [self.colors["primary"], self.colors["accent"], self.colors["accent_2"]]
        for index, item in enumerate(items):
            row, col = divmod(index, cols)
            x = mx + col * (card_w + gutter)
            y = content_y + row * (card_h + gutter)
            accent = str(item.get("accent") or palette[index % len(palette)]).lstrip("#")
            if accent in self.colors:
                accent = self.colors[accent]
            self.add_rect(
                slide,
                x,
                y,
                card_w,
                card_h,
                fill=self.colors["surface"],
                line=accent,
                rounded=False,
                name=f"CPTF_DECOMP_{slide_number}_ITEM_{index + 1}",
            )
            self.add_rect(
                slide,
                x,
                y,
                card_w,
                0.08,
                fill=accent,
                name=f"CPTF_DECOMP_{slide_number}_ACCENT_{index + 1}",
            )
            self.add_text(
                slide,
                item.get("heading", ""),
                x + 0.2,
                y + 0.25,
                card_w - 0.4,
                0.45,
                size=max(13, self.sizes["body"]),
                color=self.colors["text"],
                bold=True,
                font=self.fonts["heading"],
                name="Editable decomposition heading",
            )
            bullets = item.get("bullets")
            if isinstance(bullets, list) and bullets:
                self.add_bullets(
                    slide,
                    bullets,
                    x + 0.2,
                    y + 0.83,
                    card_w - 0.4,
                    max(0.45, card_h - 1.35),
                    size=max(10.5, self.sizes["body"] - 2),
                    name="Editable decomposition bullets",
                )
            else:
                self.add_text(
                    slide,
                    item.get("body", ""),
                    x + 0.2,
                    y + 0.83,
                    card_w - 0.4,
                    max(0.45, card_h - 1.35),
                    size=max(10.5, self.sizes["body"] - 2),
                    color=self.colors["muted_text"],
                    name="Editable decomposition body",
                )
            if str(item.get("outcome") or "").strip():
                self.add_text(
                    slide,
                    item["outcome"],
                    x + 0.2,
                    y + card_h - 0.42,
                    card_w - 0.4,
                    0.24,
                    size=max(9.5, self.sizes["small"] + 1),
                    color=accent,
                    bold=True,
                    name="Editable decomposition outcome",
                )
        self.add_takeaway(slide, data.get("takeaway"))
        self.add_footer(slide, data, slide_number)
        return slide

    def _diagram_node_style(self, node: Dict[str, Any]):
        role = str(node.get("role", "default"))
        shape_types = {
            "start": MSO_SHAPE.ROUNDED_RECTANGLE,
            "end": MSO_SHAPE.ROUNDED_RECTANGLE,
            "decision": MSO_SHAPE.DIAMOND,
            "input": getattr(MSO_SHAPE, "PARALLELOGRAM", MSO_SHAPE.RECTANGLE),
            "output": MSO_SHAPE.ROUNDED_RECTANGLE,
            "system": getattr(MSO_SHAPE, "HEXAGON", MSO_SHAPE.ROUNDED_RECTANGLE),
            "group": MSO_SHAPE.ROUNDED_RECTANGLE,
            "constraint": getattr(MSO_SHAPE, "HEXAGON", MSO_SHAPE.ROUNDED_RECTANGLE),
        }
        palette = {
            "start": (self.colors["primary"], self.colors["primary"], "FFFFFF"),
            "end": (self.colors["primary"], self.colors["primary"], "FFFFFF"),
            "decision": (self.colors["accent_2"], self.colors["accent_2"], self.colors["text"]),
            "input": (self.colors["background"], self.colors["accent"], self.colors["text"]),
            "output": (self.colors["accent"], self.colors["accent"], "FFFFFF"),
            "system": (self.colors["primary"], self.colors["primary"], "FFFFFF"),
            "group": (self.colors["surface"], self.colors["line"], self.colors["text"]),
            "constraint": (self.colors["background"], self.colors["accent_2"], self.colors["text"]),
            "process": (self.colors["surface"], self.colors["primary"], self.colors["text"]),
            "default": (self.colors["surface"], self.colors["line"], self.colors["text"]),
        }
        shape_type = shape_types.get(role, MSO_SHAPE.ROUNDED_RECTANGLE)
        fill, line, text = palette.get(role, palette["default"])
        accent = str(node.get("accent") or "").strip().lstrip("#")
        if accent in self.colors:
            fill = self.colors[accent]
        elif len(accent) == 6 and all(char in "0123456789abcdefABCDEF" for char in accent):
            fill = accent.upper()
        return shape_type, fill, line, text

    @staticmethod
    def _diagram_connection_points(
        source: Dict[str, float], target: Dict[str, float], direction: str
    ) -> tuple[float, float, float, float, bool]:
        source_cx = source["x"] + source["w"] / 2
        source_cy = source["y"] + source["h"] / 2
        target_cx = target["x"] + target["w"] / 2
        target_cy = target["y"] + target["h"] / 2
        if direction == "top_to_bottom":
            if target_cy >= source_cy:
                points = (source_cx, source["y"] + source["h"], target_cx, target["y"])
            else:
                points = (source_cx, source["y"], target_cx, target["y"] + target["h"])
            elbow = abs(source_cx - target_cx) > 0.18
        else:
            if target_cx >= source_cx:
                points = (source["x"] + source["w"], source_cy, target["x"], target_cy)
            else:
                points = (source["x"], source_cy, target["x"] + target["w"], target_cy)
            elbow = abs(source_cy - target_cy) > 0.18
        return (*points, elbow)

    def _add_diagram_node(
        self,
        slide,
        node: Dict[str, Any],
        box: Dict[str, float],
        *,
        prefix: str,
        compact: bool,
    ):
        shape_type, fill, line, text_color = self._diagram_node_style(node)
        shape = self.add_native_shape(
            slide,
            shape_type,
            box["x"],
            box["y"],
            box["w"],
            box["h"],
            fill=fill,
            line=line,
            name=f"{prefix}NODE_{node['id']}",
        )
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        margin = 0.04 if str(node.get("role")) == "decision" else 0.08
        frame.margin_left = Inches(margin)
        frame.margin_right = Inches(margin)
        frame.margin_top = Inches(0.03)
        frame.margin_bottom = Inches(0.03)
        paragraph = frame.paragraphs[0]
        label = str(node.get("label") or "")
        compact_length = len("".join(label.split()))
        label_size = 10.5 if compact else 12.5
        if compact_length > 12:
            label_size = min(label_size, 10.5)
        if compact_length > 20:
            label_size = min(label_size, 9.5)
        if str(node.get("role")) == "decision":
            label_size = min(label_size, 11.0)
        label_size = self._font_size(label_size, "diagram")
        paragraph.alignment = PP_ALIGN.CENTER
        self._format_paragraph(
            paragraph,
            label,
            size=label_size,
            color=text_color,
            font=self.fonts["body"],
            bold=True,
        )
        body = str(node.get("body") or "").strip()
        if body and box["h"] >= 0.66 and str(node.get("role")) != "decision":
            detail = frame.add_paragraph()
            detail.alignment = PP_ALIGN.CENTER
            self._format_paragraph(
                detail,
                body,
                size=self._font_size(8.5 if compact else 9.5, "diagram"),
                color=text_color,
                font=self.fonts["body"],
                bold=False,
                auto_label=bool(self.emphasis.get("enabled", True)),
            )
            detail.space_before = Pt(2)
        return shape

    def _edge_style(self, edge: Dict[str, Any]) -> tuple[str, bool]:
        edge_type = str(edge.get("type", "flow"))
        defaults = {
            "flow": (self.colors["primary"], False),
            "data": (self.colors["accent"], False),
            "control": (self.colors["accent_2"], True),
            "constraint": (self.colors["accent_2"], True),
            "feedback": (self.colors["accent"], True),
            "output": (self.colors["primary"], False),
            "reference": (self.colors["line"], True),
        }
        edge_color, dashed = defaults.get(edge_type, defaults["flow"])
        color_role = str(edge.get("color") or "").lstrip("#")
        if color_role in self.colors:
            edge_color = self.colors[color_role]
        elif len(color_role) == 6 and all(char in "0123456789abcdefABCDEF" for char in color_role):
            edge_color = color_role.upper()
        return edge_color, dashed

    def _draw_diagram(
        self,
        slide,
        diagram: Dict[str, Any],
        *,
        x: float,
        y: float,
        w: float,
        h: float,
        prefix: str,
    ) -> None:
        layout = layout_diagram(
            diagram,
            x=x,
            y=y,
            w=w,
            h=h,
        )

        for lane_index, lane in enumerate(layout["lanes"]):
            fill = self.colors["surface"] if lane_index % 2 == 0 else self.colors["background"]
            self.add_rect(
                slide,
                float(lane["x"]),
                float(lane["y"]),
                float(lane["w"]),
                float(lane["h"]),
                fill=fill,
                line=self.colors["line"],
                name=f"{prefix}LANE_{lane['id']}",
            )
            self.add_rect(
                slide,
                float(lane["x"]),
                float(lane["y"]),
                float(lane["label_w"]),
                float(lane["h"]),
                fill=self.colors["primary"],
                name=f"{prefix}LANE_LABEL_BG_{lane['id']}",
            )
            self.add_text(
                slide,
                lane["label"],
                float(lane["x"]) + 0.05,
                float(lane["y"]) + 0.05,
                float(lane["label_w"]) - 0.1,
                float(lane["h"]) - 0.1,
                size=max(9.5, self.sizes["body"] - 3),
                color="FFFFFF",
                bold=True,
                align=PP_ALIGN.CENTER,
                valign=MSO_ANCHOR.MIDDLE,
                name=f"{prefix}LANE_LABEL_{lane['id']}",
            )

        node_by_id = {str(node["id"]): node for node in diagram["nodes"]}
        direction = str(layout["direction"])
        for edge_index, edge in enumerate(layout["edges"], start=1):
            source_id = str(edge.get("from") or "")
            target_id = str(edge.get("to") or "")
            if source_id not in layout["nodes"] or target_id not in layout["nodes"]:
                continue
            x1, y1, x2, y2, elbow = self._diagram_connection_points(
                layout["nodes"][source_id], layout["nodes"][target_id], direction
            )
            edge_color, dashed = self._edge_style(edge)
            edge_type = str(edge.get("type", "flow"))
            self.add_connector(
                slide,
                x1,
                y1,
                x2,
                y2,
                elbow=elbow,
                color=edge_color,
                dashed=dashed,
                name=f"{prefix}EDGE_{edge_type.upper()}_{edge_index}_{source_id}_{target_id}",
            )
            label = str(edge.get("label") or "").strip()
            if label:
                label_w = min(0.9, max(0.42, len(label) * 0.13))
                label_shape = self.add_text(
                    slide,
                    label,
                    (x1 + x2) / 2 - label_w / 2,
                    (y1 + y2) / 2 - 0.12,
                    label_w,
                    0.24,
                    size=max(8.5, self.sizes["small"]),
                    color=edge_color,
                    bold=True,
                    align=PP_ALIGN.CENTER,
                    valign=MSO_ANCHOR.MIDDLE,
                    name=f"{prefix}EDGE_LABEL_{edge_index}",
                    margin=0.02,
                )
                label_shape.fill.solid()
                label_shape.fill.fore_color.rgb = self.rgb(self.colors["background"])
                label_shape.line.fill.background()

        compact = len(node_by_id) >= 8
        for node_id, box in layout["nodes"].items():
            node = node_by_id.get(node_id)
            if node:
                self._add_diagram_node(slide, node, box, prefix=prefix, compact=compact)

    def render_diagram(self, data: Dict[str, Any], slide_number: int):
        slide = self._new_slide()
        content_y = self.add_header(slide, data)
        diagram = data["diagram"]
        mx = float(self.geo["margin_x"])
        bottom = self.height - (1.05 if data.get("takeaway") else 0.55)
        diagram_y = content_y + 0.12
        diagram_h = max(1.0, bottom - diagram_y)
        requested_render = str(diagram.get("render", "auto"))
        prefix = (
            f"CPTF_SMARTART_{slide_number}_"
            if requested_render == "smartart"
            else f"CPTF_DIAGRAM_{slide_number}_"
        )
        self._draw_diagram(
            slide,
            diagram,
            x=mx,
            y=diagram_y,
            w=self.width - 2 * mx,
            h=diagram_h,
            prefix=prefix,
        )

        self.add_takeaway(slide, data.get("takeaway"))
        self.add_footer(slide, data, slide_number)
        return slide

    def render_composite_diagram(self, data: Dict[str, Any], slide_number: int):
        """Render two or three coordinated semantic systems on one editable slide."""

        slide = self._new_slide()
        content_y = self.add_header(slide, data)
        panels = data.get("panels", [])
        mx = float(self.geo["margin_x"])
        gutter = float(self.geo["gutter"])
        bottom = self.height - (1.05 if data.get("takeaway") else 0.55)
        panel_h = max(1.0, bottom - content_y)
        ratios = [float(panel.get("width_ratio", 1)) for panel in panels]
        usable_w = self.width - 2 * mx - gutter * (len(panels) - 1)
        ratio_total = sum(ratios)
        cursor_x = mx
        palette = [self.colors["primary"], self.colors["accent"], self.colors["accent_2"]]
        for panel_index, (panel, ratio) in enumerate(zip(panels, ratios), start=1):
            panel_w = usable_w * ratio / ratio_total
            accent = str(panel.get("accent") or palette[(panel_index - 1) % len(palette)]).lstrip("#")
            if accent in self.colors:
                accent = self.colors[accent]
            self.add_rect(
                slide,
                cursor_x,
                content_y,
                panel_w,
                panel_h,
                fill=self.colors["surface"],
                line=self.colors["line"],
                rounded=False,
                name=f"CPTF_COMPOSITE_{slide_number}_PANEL_{panel_index}",
            )
            self.add_rect(
                slide,
                cursor_x,
                content_y,
                panel_w,
                0.48,
                fill=accent,
                name=f"CPTF_COMPOSITE_{slide_number}_PANEL_HEADER_{panel_index}",
            )
            self.add_text(
                slide,
                panel.get("heading", ""),
                cursor_x + 0.12,
                content_y + 0.09,
                panel_w - 0.24,
                0.28,
                size=max(11, self.sizes["body"] - 1),
                color="FFFFFF",
                bold=True,
                align=PP_ALIGN.CENTER,
                valign=MSO_ANCHOR.MIDDLE,
                name="Editable composite panel heading",
            )
            panel_takeaway = str(panel.get("takeaway") or "").strip()
            takeaway_h = 0.42 if panel_takeaway else 0.0
            diagram_y = content_y + 0.61
            diagram_h = max(0.9, panel_h - 0.75 - takeaway_h)
            self._draw_diagram(
                slide,
                panel["diagram"],
                x=cursor_x + 0.14,
                y=diagram_y,
                w=panel_w - 0.28,
                h=diagram_h,
                prefix=f"CPTF_COMPOSITE_{slide_number}_P{panel_index}_",
            )
            if panel_takeaway:
                self.add_text(
                    slide,
                    panel_takeaway,
                    cursor_x + 0.16,
                    content_y + panel_h - 0.34,
                    panel_w - 0.32,
                    0.22,
                    size=max(8.5, self.sizes["small"]),
                    color=accent,
                    bold=True,
                    align=PP_ALIGN.CENTER,
                    name="Editable composite panel takeaway",
                )
            cursor_x += panel_w + gutter

        self.add_takeaway(slide, data.get("takeaway"))
        self.add_footer(slide, data, slide_number)
        return slide

    def render_timeline(self, data: Dict[str, Any], slide_number: int):
        slide = self._new_slide()
        content_y = self.add_header(slide, data)
        events = data.get("events", [])
        mx = float(self.geo["margin_x"])
        line_y = content_y + 2.05
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(mx + 0.35),
            Inches(line_y),
            Inches(self.width - mx - 0.35),
            Inches(line_y),
        )
        line.line.color.rgb = self.rgb(self.colors["line"])
        line.line.width = Pt(2)
        line.name = "Editable timeline line"
        step = (self.width - 2 * mx - 0.7) / max(1, len(events) - 1)
        for index, event in enumerate(events):
            center_x = mx + 0.35 + index * step
            dot = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(center_x - 0.11),
                Inches(line_y - 0.11),
                Inches(0.22),
                Inches(0.22),
            )
            dot.fill.solid()
            dot.fill.fore_color.rgb = self.rgb(self.colors["accent"] if index % 2 == 0 else self.colors["accent_2"])
            dot.line.fill.background()
            dot.name = "Editable timeline milestone"
            box_w = min(2.1, step + 0.35)
            x = max(mx, min(center_x - box_w / 2, self.width - mx - box_w))
            above = index % 2 == 0
            label_y = line_y - 1.45 if above else line_y + 0.35
            self.add_text(
                slide,
                event.get("label", ""),
                x,
                label_y,
                box_w,
                0.25,
                size=self.sizes["small"],
                color=self.colors["accent"],
                bold=True,
                align=PP_ALIGN.CENTER,
                name="Editable timeline label",
            )
            self.add_text(
                slide,
                event.get("heading", ""),
                x,
                label_y + 0.3,
                box_w,
                0.35,
                size=max(11.5, self.sizes["body"] - 1),
                color=self.colors["text"],
                bold=True,
                align=PP_ALIGN.CENTER,
                name="Editable timeline heading",
            )
            self.add_text(
                slide,
                event.get("body", ""),
                x,
                label_y + 0.72,
                box_w,
                0.55,
                size=max(9.5, self.sizes["body"] - 3),
                color=self.colors["muted_text"],
                align=PP_ALIGN.CENTER,
                name="Editable timeline body",
            )
        self.add_takeaway(slide, data.get("takeaway"))
        self.add_footer(slide, data, slide_number)
        return slide

    def render_chart(self, data: Dict[str, Any], slide_number: int):
        slide = self._new_slide()
        content_y = self.add_header(slide, data)
        chart_spec = data["chart"]
        chart_data = ChartData()
        chart_data.categories = [str(item) for item in chart_spec["categories"]]
        for series in chart_spec["series"]:
            chart_data.add_series(str(series.get("name", "Series")), series["values"])
        chart_types = {
            "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "line": XL_CHART_TYPE.LINE_MARKERS,
            "pie": XL_CHART_TYPE.PIE,
            "doughnut": XL_CHART_TYPE.DOUGHNUT,
        }
        mx = float(self.geo["margin_x"])
        chart_h = self.height - content_y - (1.15 if data.get("takeaway") else 0.65)
        frame = slide.shapes.add_chart(
            chart_types[chart_spec.get("type", "column")],
            Inches(mx),
            Inches(content_y + 0.12),
            Inches(self.width - 2 * mx),
            Inches(chart_h - 0.12),
            chart_data,
        )
        frame.name = "Editable native PowerPoint chart"
        chart = frame.chart
        chart.has_title = False
        chart.has_legend = len(chart_spec["series"]) > 1 or chart_spec.get("type") in {"pie", "doughnut"}
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.name = self.fonts["body"]
            chart.legend.font.size = Pt(self._font_size(max(9, self.sizes["small"]), "label"))
        palette = [self.colors["accent"], self.colors["primary"], self.colors["accent_2"], "8FB9A8", "B8A1D9"]
        for index, series in enumerate(chart.series):
            try:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = self.rgb(palette[index % len(palette)])
                series.format.line.color.rgb = self.rgb(palette[index % len(palette)])
            except (AttributeError, ValueError):
                pass
        if chart_spec.get("type") not in {"pie", "doughnut"}:
            chart.category_axis.tick_labels.font.name = self.fonts["body"]
            chart.category_axis.tick_labels.font.size = Pt(
                self._font_size(max(9, self.sizes["small"]), "label")
            )
            chart.value_axis.tick_labels.font.name = self.fonts["body"]
            chart.value_axis.tick_labels.font.size = Pt(
                self._font_size(max(9, self.sizes["small"]), "label")
            )
            chart.value_axis.has_major_gridlines = True
        self.add_takeaway(slide, data.get("takeaway"))
        self.add_footer(slide, data, slide_number)
        return slide

    def render_table(self, data: Dict[str, Any], slide_number: int):
        slide = self._new_slide()
        content_y = self.add_header(slide, data)
        columns = data["columns"]
        rows = data["rows"]
        mx = float(self.geo["margin_x"])
        table_h = self.height - content_y - (1.15 if data.get("takeaway") else 0.65)
        shape = slide.shapes.add_table(
            len(rows) + 1,
            len(columns),
            Inches(mx),
            Inches(content_y + 0.12),
            Inches(self.width - 2 * mx),
            Inches(table_h - 0.12),
        )
        shape.name = "Editable native PowerPoint table"
        table = shape.table
        values = [columns] + rows
        for row_index, row in enumerate(values):
            for col_index, value in enumerate(row):
                cell = table.cell(row_index, col_index)
                cell.text = str(value)
                cell.margin_left = Inches(0.08)
                cell.margin_right = Inches(0.08)
                cell.margin_top = Inches(0.05)
                cell.margin_bottom = Inches(0.05)
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.rgb(
                    self.colors["primary"] if row_index == 0 else (self.colors["surface"] if row_index % 2 else self.colors["background"])
                )
                paragraph = cell.text_frame.paragraphs[0]
                paragraph.font.name = self.fonts["body"]
                paragraph.font.size = Pt(
                    self._font_size(max(9.5, self.sizes["body"] - 2), "body")
                )
                paragraph.font.bold = row_index == 0
                paragraph.font.color.rgb = self.rgb("FFFFFF" if row_index == 0 else self.colors["text"])
                paragraph.alignment = PP_ALIGN.CENTER if row_index == 0 else PP_ALIGN.LEFT
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        self.add_takeaway(slide, data.get("takeaway"))
        self.add_footer(slide, data, slide_number)
        return slide

    def render_image_text(self, data: Dict[str, Any], slide_number: int):
        slide = self._new_slide()
        content_y = self.add_header(slide, data)
        mx = float(self.geo["margin_x"])
        gutter = float(self.geo["gutter"])
        w = (self.width - 2 * mx - gutter) / 2
        h = self.height - content_y - (1.05 if data.get("takeaway") else 0.55)
        image_left = data.get("image_side", "right") == "left"
        image_x = mx if image_left else mx + w + gutter
        text_x = mx + w + gutter if image_left else mx
        self.add_rect(slide, image_x, content_y, w, h, fill=self.colors["surface"], line=self.colors["line"], rounded=True)
        self.add_image_contain(slide, data["image"], image_x + 0.15, content_y + 0.15, w - 0.3, h - 0.3)
        self.add_bullets(slide, data.get("bullets", []), text_x, content_y + 0.12, w, h - 0.12)
        self.add_takeaway(slide, data.get("takeaway"))
        self.add_footer(slide, data, slide_number)
        return slide

    def render_closing(self, data: Dict[str, Any], slide_number: int):
        slide = self._new_slide(self.colors["primary"])
        mx = float(self.geo["margin_x"])
        self.add_text(
            slide,
            data.get("title", ""),
            mx,
            1.75,
            self.width - 2 * mx,
            1.35,
            size=self.sizes["cover_title"],
            color="FFFFFF",
            bold=True,
            font=self.fonts["heading"],
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
            name="Editable closing title",
        )
        self.add_text(
            slide,
            data.get("subtitle", ""),
            mx + 0.7,
            3.35,
            self.width - 2 * mx - 1.4,
            0.75,
            size=self.sizes["subtitle"],
            color="E8EEF3",
            align=PP_ALIGN.CENTER,
            name="Editable closing subtitle",
        )
        if data.get("contact"):
            self.add_text(
                slide,
                data["contact"],
                mx,
                5.2,
                self.width - 2 * mx,
                0.35,
                size=self.sizes["small"] + 1,
                color=self.colors["accent_2"],
                align=PP_ALIGN.CENTER,
                name="Editable contact",
            )
        if data.get("show_slide_number"):
            self.add_footer(slide, data, slide_number)
        return slide

    def render(self, data: Dict[str, Any], slide_number: int):
        slide_type = data.get("type", "title_body")
        handlers = {
            "cover": self.render_cover,
            "section": self.render_section,
            "title_body": self.render_title_body,
            "two_column": lambda item, number: self._render_columns(item, number, False),
            "comparison": lambda item, number: self._render_columns(item, number, True),
            "cards": self.render_cards,
            "decomposition": self.render_decomposition,
            "process": self.render_process,
            "diagram": self.render_diagram,
            "composite_diagram": self.render_composite_diagram,
            "timeline": self.render_timeline,
            "chart": self.render_chart,
            "table": self.render_table,
            "image_text": self.render_image_text,
            "closing": self.render_closing,
        }
        previous_style = self.style
        self._set_style(self._style_for_slide(data))
        try:
            slide = handlers[slide_type](data, slide_number)
            self._add_notes(slide, data)
            return slide
        finally:
            self._set_style(previous_style)


def _path_from_spec(spec_path: Path, raw: Any) -> Path | None:
    if not raw:
        return None
    return resolve_asset_path(spec_path, str(raw))


def _set_aspect_ratio(prs: Presentation, aspect: str) -> None:
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5 if aspect == "16:9" else 10.0)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("spec", help="Path to deck_spec.json")
    parser.add_argument("--out", help="Output PPTX path; defaults beside the spec")
    parser.add_argument("--style-profile", help="Extracted style profile JSON")
    parser.add_argument("--reference-pptx", help="Reference PPTX whose theme and masters should be retained")
    parser.add_argument("--keep-reference-slides", action="store_true", help="Append instead of removing reference slides")
    parser.add_argument(
        "--powerpoint-smartart",
        action="store_true",
        help="Use desktop PowerPoint to replace supported diagram fallbacks with native SmartArt",
    )
    args = parser.parse_args()

    spec_path = Path(args.spec).expanduser().resolve()
    try:
        spec = load_json(spec_path)
    except ValueError as exc:
        print(str(exc))
        return 2
    errors, warnings = validate_spec(spec, spec_path=spec_path, check_assets=True)
    for warning in warnings:
        print(f"WARNING: {warning}")
    if errors:
        for error in errors:
            print(f"ERROR: {error}")
        return 2

    reference = _path_from_spec(spec_path, args.reference_pptx or spec.get("reference_pptx"))
    if reference and not reference.is_file():
        print(f"Reference PPTX not found: {reference}")
        return 2
    output = Path(args.out).expanduser().resolve() if args.out else spec_path.with_name(f"{spec_path.stem}.pptx")
    if reference and output == reference:
        print("Output path must differ from the reference PPTX")
        return 2

    style_path = _path_from_spec(spec_path, args.style_profile or spec.get("style_profile"))
    if style_path:
        try:
            profile = load_json(style_path)
        except ValueError as exc:
            print(str(exc))
            return 2
    elif reference:
        profile = extract_style(reference)
    else:
        profile = None
    style = merged_style(profile, spec.get("style"))

    smartart_requested = sum(
        1
        for slide in spec["slides"]
        if isinstance(slide, dict)
        and slide.get("type") == "diagram"
        and isinstance(slide.get("diagram"), dict)
        and slide["diagram"].get("render") == "smartart"
    )
    smartart_report: Dict[str, Any] = {
        "requested": smartart_requested,
        "converted": 0,
        "mode": "native_fallback",
        "failures": [],
    }

    try:
        prs = Presentation(str(reference)) if reference else Presentation()
        initial_slides = len(prs.slides)
        if reference and not args.keep_reference_slides:
            _remove_all_slides(prs)
            initial_slides = 0
        if not reference or spec.get("force_aspect_ratio"):
            _set_aspect_ratio(prs, spec.get("aspect_ratio", "16:9"))
        prs.core_properties.title = str(spec.get("title") or spec_path.stem)
        prs.core_properties.subject = "Native editable deck generated by codex-ppt-simple"
        prs.core_properties.comments = "Text, shapes, charts, tables, and diagrams are native PowerPoint objects."
        builder = DeckBuilder(prs, spec, style, spec_path)
        for index, slide_data in enumerate(spec["slides"], start=1):
            builder.render(slide_data, initial_slides + index)
        output.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output))
        if args.powerpoint_smartart and smartart_requested:
            from apply_smartart import apply_smartart

            converted = apply_smartart(output, spec_path, slide_offset=initial_slides)
            smartart_report = {
                "requested": smartart_requested,
                "mode": "powerpoint_native",
                **converted,
            }
    except Exception as exc:
        print(f"Build failed: {exc}")
        return 1

    summary = {
        "output": str(output),
        "generated_slides": len(spec["slides"]),
        "total_slides": len(prs.slides),
        "mode": spec.get("mode", "hybrid"),
        "reference_pptx": str(reference) if reference else None,
        "style_profile": str(style_path) if style_path else ("auto-extracted" if reference else "default"),
        "smartart": smartart_report,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
